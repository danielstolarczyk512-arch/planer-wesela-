"""Syntetyczny korpus ofert weselnych do testów ekstrakcji (wyłącznie fikcyjne dane).

20 dokumentów: 10 PDF z warstwą tekstową, 6 DOCX, 4 skany (PNG, JPG, 2x PDF bez tekstu).
Każdy ma listę fraz, które MUSZĄ znaleźć się w wyodrębnionym tekście.
"""

import io
import shutil
import unicodedata
from dataclasses import dataclass, field
from functools import cache
from pathlib import Path

FONT_CANDIDATES = (
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",  # Debian (obraz Dockera)
    "C:/Windows/Fonts/arial.ttf",  # Windows
    "/Library/Fonts/Arial.ttf",
)


@dataclass(frozen=True)
class Sample:
    name: str
    data: bytes
    mime: str
    expected: list[str]
    needs_ocr: bool = False
    lines: list[str] = field(default_factory=list)


VENUES = [
    ("Dworek Pod Lipami", 330, 80, 5000, 150),
    ("Sala Bankietowa Złoty Róg", 289, 100, 4000, 120),
    ("Pałac Różany Ogród", 410, 60, 8000, 180),
    ("Karczma Pod Jeleniem", 250, 70, 3000, 110),
    ("Hotel Nad Jeziorem", 360, 90, 6000, 160),
    ("Stodoła Wiśniowy Sad", 300, 50, 4500, 140),
    ("Zajazd Łąkowy Zakątek", 275, 80, 3500, 125),
    ("Folwark Żurawie Gniazdo", 395, 70, 7000, 170),
    ("Restauracja Śródmiejska", 265, 60, 3000, 115),
    ("Oranżeria Biała Róża", 420, 100, 9000, 190),
]
SERVICES = [
    ("Zespół Weselny Krakowiacy", "Oprawa muzyczna 8 godzin", 7500, 1500),
    ("DJ Mateusz Muzyka", "Muzyka do białego rana", 4200, 1000),
    ("Fotograf Ślubny Ania Kadr", "Reportaż ślubny i sesja plenerowa", 6800, 2000),
    ("Studio Filmowe Żółty Kadr", "Film ślubny 4K z dronem", 7900, 2500),
    ("Kwiaciarnia Łąka i Róża", "Dekoracja kościoła i sali", 5200, 1000),
    ("Słodki Stół Bezowa Chmurka", "Tort trzypiętrowy i słodki stół", 3100, 500),
]


def offer_lines(name: str, price: int, minimum: int, deposit: int, after: int) -> list[str]:
    return [
        f"Oferta weselna: {name}",
        f"Menu weselne: {price} zł za osobę dorosłą",
        "Dzieci 4-10 lat: 50% ceny",
        f"Minimalna liczba gości: {minimum} osób",
        f"Poprawiny: {after} zł za osobę",
        f"Zadatek: {deposit} zł płatny przy podpisaniu umowy",
        "Napoje gazowane i woda w cenie. Alkohol własny, korkowe 0 zł.",
        "Odwołanie przyjęcia: zadatek przepada.",
    ]


@cache
def font_path() -> str | None:
    return next((p for p in FONT_CANDIDATES if Path(p).exists()), None)


def _pdf_text(lines: list[str]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    font = font_path()
    if font is None:
        raise RuntimeError("Brak czcionki TTF z polskimi znakami do generowania PDF")
    pdfmetrics.registerFont(TTFont("PL", font))
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setFont("PL", 12)
    y = 800
    for i, line in enumerate(lines):
        if i == 5:  # druga strona: test podziału na strony
            c.showPage()
            c.setFont("PL", 12)
            y = 800
        c.drawString(60, y, line)
        y -= 22
    c.save()
    return buf.getvalue()


def _docx(lines: list[str], table: list[tuple[str, str]]) -> bytes:
    from docx import Document

    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    t = doc.add_table(rows=0, cols=2)
    for a, b in table:
        row = t.add_row().cells
        row[0].text, row[1].text = a, b
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _image(lines: list[str], fmt: str) -> bytes:
    from PIL import Image, ImageDraw, ImageFont

    font = font_path()
    if font is None:
        raise RuntimeError("Brak czcionki TTF do generowania skanów")
    img = Image.new("RGB", (1700, 120 + 70 * len(lines)), "white")
    draw = ImageDraw.Draw(img)
    face = ImageFont.truetype(font, 44)
    for i, line in enumerate(lines):
        draw.text((60, 60 + 70 * i), line, fill="black", font=face)
    buf = io.BytesIO()
    img.save(buf, format=fmt, **({"quality": 95} if fmt == "JPEG" else {}))
    return buf.getvalue()


@cache
def corpus() -> tuple[Sample, ...]:
    samples: list[Sample] = []
    for name, price, minimum, deposit, after in VENUES:  # 10 x PDF z tekstem
        lines = offer_lines(name, price, minimum, deposit, after)
        samples.append(
            Sample(
                name=f"pdf-{name}",
                data=_pdf_text(lines),
                mime="application/pdf",
                lines=lines,
                expected=[
                    name,
                    f"{price} zł za osobę",
                    f"{minimum} osób",
                    f"Zadatek: {deposit} zł",
                    f"Poprawiny: {after} zł",
                ],
            )
        )
    for vendor, desc, price, deposit in SERVICES:  # 6 x DOCX
        lines = [
            f"Oferta: {vendor}",
            desc,
            f"Cena pakietu: {price} zł",
            f"Zaliczka {deposit} zł, dopłata 7 dni przed ślubem",
        ]
        data = _docx(lines, [("Pakiet", f"{price} zł"), ("Nadgodzina", "400 zł")])
        samples.append(
            Sample(
                name=f"docx-{vendor}",
                lines=lines,
                mime=("application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
                data=data,
                expected=[
                    vendor,
                    desc,
                    f"{price} zł",
                    f"Zaliczka {deposit} zł",
                    "Nadgodzina | 400 zł",
                ],
            )
        )
    for i, (name, price, minimum, deposit, _after) in enumerate(VENUES[:4]):  # 4 x skan
        lines = [
            f"Oferta {name}",
            f"Menu {price} zl za osobe",
            f"Minimum {minimum} osob",
            f"Zadatek {deposit} zl",
        ]
        fmt = ("PNG", "JPEG", "PDF", "PDF")[i]
        mime = {"PNG": "image/png", "JPEG": "image/jpeg", "PDF": "application/pdf"}[fmt]
        samples.append(
            Sample(
                name=f"scan-{fmt.lower()}-{i}",
                data=_image(lines, fmt),
                mime=mime,
                lines=lines,
                expected=[str(price), str(minimum), str(deposit), "Zadatek"],
                needs_ocr=True,
            )
        )
    return tuple(samples)


def tesseract_available() -> bool:
    return shutil.which("tesseract") is not None


def normalize(text: str, strip_diacritics: bool = False) -> str:
    text = " ".join(text.split()).lower()
    if strip_diacritics:
        text = text.replace("ł", "l")
        text = "".join(
            c for c in unicodedata.normalize("NFKD", text) if not unicodedata.combining(c)
        )
    return text
